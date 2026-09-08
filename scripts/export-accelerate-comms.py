#!/usr/bin/env python3
"""Export the seven reviewed Accelerate Markdown bodies to Word, outside Git."""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

NAMES = ['master-links', 'vis-rsm-email', 'aes-rsm-email',
         'vis-prospect-email-flight', 'vis-prospect-email-non-flight',
         'aes-prospect-email-flight', 'aes-prospect-email-non-flight']


def add_text(paragraph, text):
    # Templates use labeled bare URLs. Preserve them as Word hyperlinks.
    offset = 0
    for match in re.finditer(r'https?://[^\s<>]+', text):
        paragraph.add_run(text[offset:match.start()])
        url = match.group().rstrip('.,;')
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
        props.append(style); run.append(props)
        node = OxmlElement('w:t'); node.text = url
        run.append(node); link.append(run); paragraph._p.append(link)
        paragraph.add_run(match.group()[len(url):])
        offset = match.end()
    paragraph.add_run(text[offset:])


def export(text, destination):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size in [('Title', 18), ('Heading 1', 13), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(5)
    for block in re.split(r'\n\s*\n', text.strip()):
        if block.startswith('# '):
            p = doc.add_paragraph(style='Title'); add_text(p, block[2:])
        elif block.startswith('## '):
            p = doc.add_paragraph(style='Heading 1'); add_text(p, block[3:])
        elif block.startswith('### '):
            p = doc.add_paragraph(style='Heading 2'); add_text(p, block[4:])
        elif all(line.startswith('- ') for line in block.splitlines()):
            for line in block.splitlines(): add_text(doc.add_paragraph(style='List Bullet'), line[2:])
        else:
            add_text(doc.add_paragraph(), block)
    for p in doc.paragraphs:
        if p.text.strip() in ['Best,', 'Best regards,']:
            p.paragraph_format.keep_with_next = True
    doc.save(destination)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', required=True, type=Path)
    parser.add_argument('--output-dir', required=True, type=Path)
    parser.add_argument('--city-slug', required=True)
    args = parser.parse_args()
    if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', args.city_slug):
        parser.error('city-slug must be lowercase words separated by hyphens')
    repo = Path(__file__).resolve().parents[1]
    for folder in [args.input_dir.resolve(), args.output_dir.resolve()]:
        if folder == repo or repo in folder.parents:
            parser.error('Use external working and output folders, outside this repository')
    bodies = {}
    for name in NAMES:
        source = args.input_dir / (name + '.md')
        if not source.is_file(): parser.error('Missing required file: ' + source.name)
        text = source.read_text(encoding='utf-8')
        if re.search(r'\{\{[^}]+\}\}|\[(?:PLACEHOLDER|Name|Rep Name)[^\]]*\]', text, re.I):
            parser.error('Unresolved field in ' + source.name)
        destination = args.output_dir / (args.city_slug + '-' + name + '.docx')
        if destination.exists(): parser.error('Refusing to overwrite ' + destination.name)
        bodies[destination] = text
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for destination, text in bodies.items():
        export(text, destination)
        print(destination.name)


if __name__ == '__main__':
    main()
